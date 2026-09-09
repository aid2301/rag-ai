"""XLSX 解析（每个工作表转为 Markdown 表格，作为独立章节）。"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from app.services.document_parser.base import ParsedDocument, ParsedSection


def _sheet_to_markdown(rows: list[list[str]], max_rows: int = 500) -> str:
    if not rows:
        return ""
    header = rows[0]
    lines = ["| " + " | ".join(_cell(c) for c in header) + " |"]
    lines.append("| " + " | ".join(["---"] * len(header)) + " |")
    for row in rows[1:max_rows]:
        cells = [_cell(c) for c in row]
        # 补齐列数
        while len(cells) < len(header):
            cells.append("")
        lines.append("| " + " | ".join(cells[: len(header)]) + " |")
    return "\n".join(lines)


def _cell(v) -> str:
    if v is None:
        return ""
    return str(v).replace("\n", " ").replace("|", "\\|").strip()


def parse_xlsx(path: str, filename: str) -> ParsedDocument:
    wb = load_workbook(path, read_only=True, data_only=True)
    title = Path(filename).stem
    sections: list[ParsedSection] = []
    full_parts: list[str] = []

    for idx, ws in enumerate(wb.worksheets, start=1):
        rows = [[_cell(c) for c in row] for row in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(x for x in r)]
        if not rows:
            continue
        md = _sheet_to_markdown(rows)
        heading = ws.title or f"Sheet{idx}"
        full_parts.append(f"# {heading}\n\n{md}")
        sections.append(
            ParsedSection(heading=heading, content=md, level=1, page_number=None, parent_heading=None)
        )

    wb.close()
    if not sections:
        raise ValueError("XLSX 中没有可解析的数据")

    full_text = "\n\n".join(full_parts)
    return ParsedDocument(title=title, full_text=full_text, sections=sections)
