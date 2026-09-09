"""DOCX 解析（保留标题层级与表格，按文档顺序）。"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.document import Document as _Doc
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

from app.services.document_parser.base import ParsedDocument, split_into_sections

_HEADING_MAP = {
    "heading 1": 1,
    "heading 2": 2,
    "heading 3": 3,
    "heading 4": 4,
    "heading 5": 5,
    "heading 6": 6,
    "标题 1": 1,
    "标题 2": 2,
    "标题 3": 3,
    "标题 4": 4,
}


def _heading_level(style_name: str) -> int | None:
    name = (style_name or "").lower()
    for key, level in _HEADING_MAP.items():
        if name == key or name.startswith(key + " ") or name.startswith(key):
            return level
    return None


def _cell_text(cell: _Cell) -> str:
    return " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())


def _table_to_markdown(table: Table) -> str:
    rows = table.rows
    if not rows:
        return ""
    lines: list[str] = []
    header = [_cell_text(c) for c in rows[0].cells]
    lines.append("| " + " | ".join(header) + " |")
    lines.append("| " + " | ".join(["---"] * len(header)) + " |")
    for row in rows[1:]:
        cells = [_cell_text(c) for c in row.cells]
        lines.append("| " + " | ".join(cells) + " |")
    return "\n".join(lines)


def parse_docx(path: str, filename: str) -> ParsedDocument:
    doc: _Doc = Document(path)
    title = Path(filename).stem

    # 按文档顺序遍历段落与表格
    lines: list[tuple[str, int | None]] = []
    full_parts: list[str] = []
    body = doc.element.body

    from docx.text.paragraph import Paragraph as _P
    from docx.table import Table as _T

    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            para = _P(child, doc)
            text = para.text
            style = para.style.name if para.style else ""
            level = _heading_level(style)
            if text.strip():
                if level is not None:
                    # 转成 Markdown 标题，让 split_into_sections 正确识别层级
                    md_heading = f"{'#' * level} {text}"
                    full_parts.append(md_heading)
                    lines.append((md_heading, None))
                else:
                    full_parts.append(text)
                    lines.append((text, None))
        elif child.tag == qn("w:tbl"):
            table = _T(child, doc)
            md = _table_to_markdown(table)
            if md:
                full_parts.append(md)
                for line in md.splitlines():
                    lines.append((line, None))

    full_text = "\n".join(full_parts)
    if not full_text.strip():
        raise ValueError("DOCX 文档内容为空")

    sections = split_into_sections(lines, title)
    return ParsedDocument(title=title, full_text=full_text, sections=sections)
